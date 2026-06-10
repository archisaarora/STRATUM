"""Tests for the alternate access paths added after real-world recon:
SIPRI TIV tables, Comtrade public preview, World Bank Data360."""
import io
import json

import openpyxl
import pandas as pd

from stratum.core.clients import comtrade, worldbank
from stratum.core.features.arms_velocity import transfer_velocity
from stratum.core.parsing.sipri_arms import parse_tiv_table
from stratum.core.scoring.signals import arms_transfer_spike_signals


def _tiv_table_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TIV imports"
    ws.append(["TIV of arms imports, SAMPLE", None, None, None, None, None])
    ws.append([None, None, None, None, None, None])
    ws.append(["Recipient", 2020, 2021, 2022, 2023, 2024])
    ws.append(["Myanmar", 50, 50, 60, 55, 400])
    ws.append(["Pakistan", 300, 310, 290, 305, 300])
    ws.append(["France", 80, 0, 90, 85, 80])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_tiv_table_shape():
    df = parse_tiv_table(_tiv_table_xlsx())
    assert set(df["status"]) == {"tiv_annual"}
    mmr = df[df["recipient"] == "Myanmar"]
    assert len(mmr) == 5
    assert mmr[mmr["delivery_year_last"] == 2024]["total_tiv"].iloc[0] == 400
    # zero-TIV cells dropped (France 2021)
    fra = df[df["recipient"] == "France"]
    assert 2021 not in set(fra["delivery_year_last"])
    assert df["transfer_id"].is_unique


def test_velocity_and_signal_from_tiv_table():
    df = parse_tiv_table(_tiv_table_xlsx())
    df = df.rename(columns={"recipient": "recipient_country"})
    df["recipient_country"] = df["recipient_country"].map(
        {"Myanmar": "MMR", "Pakistan": "PAK", "France": "FRA"})
    vel = transfer_velocity(df)
    mmr_2024 = vel[(vel.recipient_country == "MMR") & (vel.year == 2024)].iloc[0]
    assert mmr_2024["transfer_spike_flag"]
    assert mmr_2024["top_category"] is None  # no weapon detail in TIV tables

    signals = arms_transfer_spike_signals(vel)
    mmr_sig = [s for s in signals if s["country_code"] == "MMR"]
    assert len(mmr_sig) == 1
    assert "mixed categories" in mmr_sig[0]["description_text"]


def test_public_work_queue_grain_and_resume():
    queue = comtrade.build_public_work_queue(
        {"IRN": 364, "CHN": 156}, [2023, 2024], ["3601", "8108"],
        done={("IRN", 2023, "3601")})
    assert ("IRN", 2023, "3601") not in queue
    assert ("CHN", 2023, "3601") in queue
    assert len(queue) == 2 * 2 * 2 - 1


class FakeResponse:
    def __init__(self, payload):
        self.payload = payload
        self.status_code = 200
        self.headers = {}

    def json(self):
        return self.payload

    def raise_for_status(self):
        pass


class FakeSession:
    """Captures request params and replays canned payloads."""

    def __init__(self, payloads):
        self.payloads = list(payloads)
        self.calls = []

    def request(self, method, url, **kwargs):
        self.calls.append((method, url, kwargs))
        return FakeResponse(self.payloads.pop(0))


def test_public_preview_params_match_spec():
    session = FakeSession([{"count": 1, "data": [
        {"reporterCode": 364, "cmdCode": "3601", "flowCode": "M",
         "primaryValue": 1.0, "refYear": 2023}]}])
    recs = comtrade.fetch_public_preview(
        session, "https://comtradeapi.un.org", reporter_m49=364,
        year=2023, hs_code="3601", throttle_seconds=0)
    assert len(recs) == 1
    method, url, kwargs = session.calls[0]
    assert url.endswith("/public/v1/preview/C/A/HS")
    params = kwargs["params"]
    assert params["cmdCode"] == "3601"        # exactly one code (spec limit)
    assert params["period"] == 2023           # exactly one period
    assert params["partnerCode"] == 0         # world aggregate
    assert params["flowCode"] == "M,X"
    assert "subscription-key" not in params   # keyless endpoint


def test_data360_mapping_and_pagination():
    page1 = {"count": 1001, "value": [
        {"OBS_VALUE": "3.43", "REF_AREA": "USA", "TIME_PERIOD": "2023",
         "INDICATOR": "WB_WDI_MS_MIL_XPND_GD_ZS"}] * 1000}
    page2 = {"count": 1001, "value": [
        {"OBS_VALUE": None, "REF_AREA": "IRN", "TIME_PERIOD": "2024",
         "INDICATOR": "WB_WDI_MS_MIL_XPND_GD_ZS"}]}
    session = FakeSession([page1, page2])
    rows = worldbank.fetch_indicator_data360(
        session, "https://data360api.worldbank.org", "MS.MIL.XPND.GD.ZS",
        start_year=2020, end_year=2025)
    assert len(rows) == 1001
    assert session.calls[0][2]["params"]["INDICATOR"] == "WB_WDI_MS_MIL_XPND_GD_ZS"
    assert session.calls[1][2]["params"]["skip"] == 1000
    assert rows[0] == {"indicator_code": "MS.MIL.XPND.GD.ZS",
                       "country_iso3": "USA", "country_name": None,
                       "year": 2023, "value": 3.43}
    assert rows[-1]["value"] is None  # null OBS_VALUE survives as None


def test_sipri_pct_gdp_fraction_scaling():
    """Real SIPRI files store share-of-GDP as fractions (0.034 = 3.4%)."""
    from stratum.core.parsing.sipri_milex import parse_workbook

    wb = openpyxl.Workbook()
    cur = wb.active
    cur.title = "Current US$"
    cur.append(["Country", "Notes", 2023, 2024])
    cur.append(["Iran", "", 10_000, 10_400])
    share = wb.create_sheet("Share of GDP")
    share.append(["Country", "Notes", 2023, 2024])
    share.append(["Iran", "", 0.022, 0.0222])
    share.append(["Russia", "", 0.059, 0.0687])
    buf = io.BytesIO()
    wb.save(buf)

    df = parse_workbook(buf.getvalue())
    pct = df[df.measure == "expenditure_pct_gdp"]
    rus = pct[(pct.country_name_raw == "Russia") & (pct.year == 2024)]
    assert abs(rus.iloc[0]["value"] - 6.87) < 0.01  # scaled to percent
